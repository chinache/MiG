import argparse
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


TYPE_MAP = {
    '单选': 'single',
    '单选题': 'single',
    '多选': 'multiple',
    '多选题': 'multiple',
    '判断': 'judge',
    '判断题': 'judge',
}
TYPE_CODE = {'single': 'S', 'multiple': 'M', 'judge': 'J'}
OPTION_KEYS = ('A', 'B', 'C', 'D', 'E')
TRUE_ANSWERS = {'对', '正确', '是', '√', 'true', 'yes'}
FALSE_ANSWERS = {'错', '错误', '否', '×', 'false', 'no'}


def clean_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\s+', ' ', str(value)).strip()


def decode_zip_name(name):
    try:
        return name.encode('cp437').decode('gbk')
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name


def normalize_header(value):
    return re.sub(r'\s+', '', clean_text(value))


def find_header_row(worksheet):
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 12), values_only=True),
        start=1,
    ):
        headers = [normalize_header(value) for value in row]
        if '题干' in headers and '正确答案' in headers:
            return row_number, headers
    raise ValueError(f'{worksheet.title} 未找到包含“题干”和“正确答案”的表头')


def bank_and_category(file_stem, title):
    bank = re.split(r'[·•]', clean_text(title), maxsplit=1)[0].strip()
    if not bank:
        bank = re.sub(r'题库$', '', file_stem).strip()

    if bank.startswith('合规天天练（') and bank.endswith('）'):
        category = bank[len('合规天天练（'):-1].strip()
    elif bank.startswith('合规天天练(') and bank.endswith(')'):
        category = bank[len('合规天天练('):-1].strip()
    else:
        category = re.sub(r'^合规天天练\s*[-—－]?\s*', '', bank)
        category = re.sub(r'题库$', '', category).strip()
    return bank, category or '综合'


def question_type(raw_type, sheet_name):
    normalized = clean_text(raw_type)
    normalized_sheet = clean_text(sheet_name)
    result = TYPE_MAP.get(normalized) or TYPE_MAP.get(normalized_sheet)
    if not result:
        raise ValueError(f'无法识别题型：{normalized or normalized_sheet}')
    return result


def normalize_answer(raw_answer, item_type):
    answer_text = clean_text(raw_answer)
    lowered = answer_text.lower()
    if item_type == 'judge':
        if answer_text in TRUE_ANSWERS or lowered in TRUE_ANSWERS:
            return ['A']
        if answer_text in FALSE_ANSWERS or lowered in FALSE_ANSWERS:
            return ['B']

    answers = []
    for key in re.findall(r'[A-E]', answer_text.upper()):
        if key not in answers:
            answers.append(key)
    if not answers:
        raise ValueError(f'无法识别正确答案：{answer_text!r}')
    return answers


def make_id(file_name, sheet_name, source_number, stem):
    source_key = '|'.join([file_name, sheet_name, source_number, stem])
    return 'daily-' + hashlib.sha1(source_key.encode('utf-8')).hexdigest()[:16]


def convert_archive(zip_path):
    questions = []
    reports = []
    groups = []

    with zipfile.ZipFile(zip_path) as archive:
        workbook_members = [
            info for info in archive.infolist()
            if not info.is_dir() and decode_zip_name(info.filename).lower().endswith('.xlsx')
        ]
        workbook_members.sort(key=lambda info: decode_zip_name(info.filename))

        for workbook_index, info in enumerate(workbook_members, start=1):
            file_name = Path(decode_zip_name(info.filename)).name
            workbook = load_workbook(BytesIO(archive.read(info)), read_only=True, data_only=False)
            workbook_count = 0

            for worksheet in workbook.worksheets:
                header_row, headers = find_header_row(worksheet)
                header_index = {header: index for index, header in enumerate(headers) if header}
                title = worksheet.cell(row=1, column=1).value
                bank, category = bank_and_category(Path(file_name).stem, title)
                sheet_count = 0
                sheet_questions = []

                for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                    def field(name):
                        index = header_index.get(name)
                        return row[index] if index is not None and index < len(row) else None

                    stem = clean_text(field('题干'))
                    if not stem:
                        continue

                    source_number = clean_text(field('序号')) or str(sheet_count + 1)
                    item_type = question_type(field('题型'), worksheet.title)
                    options = {}
                    if item_type == 'judge':
                        options = {'A': '正确', 'B': '错误'}
                    else:
                        for key in OPTION_KEYS:
                            option_text = clean_text(field('选项' + key))
                            if option_text:
                                options[key] = option_text

                    answer = normalize_answer(field('正确答案'), item_type)
                    missing_options = [key for key in answer if key not in options]
                    if missing_options:
                        raise ValueError(
                            f'{file_name}/{worksheet.title}/{source_number} '
                            f'答案 {missing_options} 没有对应选项'
                        )

                    number = f'{category}-{TYPE_CODE[item_type]}-{int(float(source_number)):03d}' \
                        if re.fullmatch(r'\d+(?:\.0+)?', source_number) \
                        else f'{category}-{TYPE_CODE[item_type]}-{source_number}'
                    question = {
                        'id': make_id(file_name, worksheet.title, source_number, stem),
                        'type': item_type,
                        'stem': stem,
                        'options': options,
                        'answer': answer,
                        'analysis': clean_text(field('答案解析')),
                        'bank': bank,
                        'category': category,
                        'knowledgePoints': [category],
                        'number': number,
                    }
                    questions.append(question)
                    sheet_questions.append(question)
                    sheet_count += 1

                workbook_count += sheet_count
                reports.append({
                    'file': file_name,
                    'sheet': worksheet.title,
                    'bank': bank,
                    'category': category,
                    'questions': sheet_count,
                })
                groups.append({
                    'fileName': f'daily-{workbook_index:02d}-{TYPE_CODE[question_type(None, worksheet.title)].lower()}.json',
                    'questions': sheet_questions,
                })

            workbook.close()
            if workbook_count == 0:
                raise ValueError(f'{file_name} 未读取到题目')

    ids = [question['id'] for question in questions]
    if len(ids) != len(set(ids)):
        raise ValueError('转换结果中存在重复 id')

    return questions, reports, groups


def main():
    parser = argparse.ArgumentParser(description='将“合规天天练”Excel 压缩包转换为知心题库 JSON')
    parser.add_argument('zip_path', type=Path)
    parser.add_argument('--output-dir', type=Path)
    args = parser.parse_args()

    questions, reports, groups = convert_archive(args.zip_path)
    type_counts = Counter(question['type'] for question in questions)
    bank_counts = Counter(question['bank'] for question in questions)
    duplicate_stems = sum(count - 1 for count in Counter(question['stem'] for question in questions).values() if count > 1)
    json_bytes = 0
    if args.output_dir:
        args.output_dir.mkdir(parents=True, exist_ok=True)
        for old_file in args.output_dir.glob('daily-*.json'):
            old_file.unlink()
        for group in groups:
            payload = json.dumps(group['questions'], ensure_ascii=False, indent=2) + '\n'
            (args.output_dir / group['fileName']).write_text(payload, encoding='utf-8')
            json_bytes += len(payload.encode('utf-8'))
    else:
        json_bytes = len((json.dumps(questions, ensure_ascii=False, indent=2) + '\n').encode('utf-8'))

    result = {
        'workbooks': len({report['file'] for report in reports}),
        'sheets': len(reports),
        'questions': len(questions),
        'types': dict(type_counts),
        'banks': dict(bank_counts),
        'duplicateStemOccurrences': duplicate_stems,
        'jsonFiles': len(groups),
        'jsonBytes': json_bytes,
        'outputDirectory': str(args.output_dir) if args.output_dir else None,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(f'导入失败：{error}', file=sys.stderr)
        raise
